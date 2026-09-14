import os
import fitz
from pptx import Presentation
from docx import Document


def extract_text_from_pdf(file_path):
    document = fitz.open(file_path)

    text = ""

    for page in document:
        text += page.get_text() + "\n"

    document.close()

    return text


def extract_text_from_pptx(file_path):
    """Extract text from a PowerPoint file."""
    presentation = Presentation(file_path)

    text = ""

    for slide_number, slide in enumerate(presentation.slides, start=1):
        text += f"\n--- Slide {slide_number} ---\n"

        for shape in slide.shapes:

            if hasattr(shape, "text"):
                if shape.text.strip():
                    text += shape.text + "\n"

    return text


def extract_text_from_docx(file_path):
    """Extract text from a Word document."""
    document = Document(file_path)

    text = ""

    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            text += paragraph.text + "\n"

    # Extract text from tables also
    for table in document.tables:

        text += "\n--- Table ---\n"

        for row in table.rows:

            row_text = []

            for cell in row.cells:
                row_text.append(cell.text.strip())

            text += " | ".join(row_text) + "\n"

    return text


def extract_text(file_path):
    """
    Automatically detect the document type
    and extract its text.
    """

    if not os.path.exists(file_path):
        raise FileNotFoundError(
            f"File not found: {file_path}"
        )

    extension = os.path.splitext(file_path)[1].lower()

    if extension == ".pdf":

        return extract_text_from_pdf(file_path)

    elif extension == ".pptx":

        return extract_text_from_pptx(file_path)

    elif extension == ".docx":

        return extract_text_from_docx(file_path)

    else:

        raise ValueError(
            "Unsupported file format. "
            "Supported formats: PDF, PPTX, DOCX"
        )


if __name__ == "__main__":

    print("\n===================================")
    print("     TECH SPARK DOCUMENT TEST")
    print("===================================\n")

    file_path = input(
        "Enter document path: "
    ).strip()

    try:

        extracted_text = extract_text(file_path)

        print("\nDocument processed successfully!")
        print("-----------------------------------")

        print(
            "\nExtracted text preview:\n"
        )

        print(extracted_text[:2000])

        print("\n-----------------------------------")

        print(
            f"Total characters extracted: "
            f"{len(extracted_text)}"
        )

    except Exception as error:

        print("\nERROR:")
        print(error)